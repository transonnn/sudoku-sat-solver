import csv
import statistics
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_word_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_code

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "Open in Word to update this field"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"
REPORT_DIR = ROOT / "report"
OUTPUT = REPORT_DIR / "sudoku_solver_report.docx"

REPORT_DIR.mkdir(parents=True, exist_ok=True)


def read_csv(path):
    if not path.exists():
        print(f"WARNING: Missing CSV: {path}")
        return []
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def set_cell_text(cell, text, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(9)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)

    doc.add_paragraph()
    return table


def add_figure(doc, path, caption, width=6.0):
    if not path.exists():
        doc.add_paragraph(f"[Figure unavailable: {path.name}]")
        return

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))

    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(9)


def fmt_seconds(value):
    return f"{value:.4f}"


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


doc = Document()

# Document formatting
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)

for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Times New Roman"

styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)

# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n\n")

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("COMPARATIVE STUDY OF SUDOKU SOLVERS")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(20)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run(
    "\nSAT, CP-SAT, Mixed-Integer Programming, "
    "and Scaled Sudoku Instances"
)
r.font.size = Pt(14)

doc.add_paragraph("\n\n")
p = doc.add_paragraph("Research Project Report")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(
    "\nAuthor: [Add author name]\n"
    "Institution: [Add institution]\n"
    "Course: [Add course name]\n"
    "Date: [Add submission date]"
)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# Table of contents
doc.add_heading("Table of Contents", level=1)
toc_paragraph = doc.add_paragraph()
add_word_field(toc_paragraph, 'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()

# Abstract
doc.add_heading("Abstract", level=1)
doc.add_paragraph(
    "This project studies several approaches to solving Sudoku as a "
    "constraint satisfaction problem. The implemented methods include "
    "Boolean satisfiability (SAT), Google OR-Tools CP-SAT, IBM CPLEX "
    "mixed-integer programming (MIP), and Gurobi MIP. The evaluation "
    "contains a 9x9 dataset of 30 instances grouped by difficulty and "
    "additional scaled Sudoku instances of sizes 16x16 and 25x25. "
    "The report summarizes measured solving times and successful "
    "solution checks from the stored benchmark results. Performance "
    "conclusions are restricted to these datasets and the tested "
    "software and hardware configuration."
)
doc.add_paragraph(
    "Keywords: Sudoku, SAT, cardinality constraints, CP-SAT, MIP, "
    "benchmarking, constraint programming."
)

# 1
doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "Sudoku is a constraint satisfaction problem in which symbols must "
    "be assigned to cells while satisfying row, column, and subgrid "
    "constraints. Although the rules are simple, the search space grows "
    "rapidly with grid size. Sudoku therefore provides a useful setting "
    "for comparing general-purpose constraint and optimization solvers."
)
doc.add_paragraph(
    "The goal of this project is to implement multiple mathematical "
    "formulations, validate their solutions, and compare empirical "
    "solving times. The project also investigates SAT cardinality "
    "encodings and extends the evaluation beyond the standard 9x9 grid."
)

doc.add_heading("1.1 Research questions", level=2)
add_bullet(
    doc, "How do SAT, CP-SAT, CPLEX MIP, and Gurobi MIP compare on 9x9 Sudoku?"
)
add_bullet(
    doc, "How does solving time change for 16x16 and 25x25 instances?"
)
add_bullet(
    doc, "What limitations must be considered when interpreting benchmark results?"
)

# 2
doc.add_heading("2. Mathematical formulation", level=1)
doc.add_heading("2.1 Generalized Sudoku", level=2)
doc.add_paragraph(
    "Let the grid have size N x N, where N = q^2 for an integer q. "
    "Each cell (r,c) receives a symbol from {1,...,N}. Every row, "
    "column, and q x q subgrid must contain each symbol exactly once. "
    "Given clues are fixed to their original values."
)
doc.add_paragraph(
    "The number of possible assignments before constraints are applied "
    "is N^(N^2). The constraints eliminate invalid assignments, but "
    "the search difficulty depends strongly on the instance and formulation."
)

doc.add_heading("2.2 SAT formulation", level=2)
doc.add_paragraph(
    "Define a Boolean variable x[r,c,v] that is true if cell (r,c) "
    "contains value v. There are N^3 such variables. A valid encoding "
    "enforces that each cell has exactly one value, each value occurs "
    "exactly once in every row and column, and each value occurs exactly "
    "once in each subgrid. Clue constraints fix the corresponding variables."
)
doc.add_paragraph(
    "For a set of Boolean variables X, an exactly-one constraint can be "
    "written as an at-least-one clause together with an at-most-one "
    "constraint:"
)
doc.add_paragraph("OR(x in X) AND AMO(X).")
doc.add_paragraph(
    "The pairwise AMO encoding adds a clause (not x_i OR not x_j) for "
    "every distinct pair i < j. It is simple but uses quadratically many "
    "binary clauses in the size of X. Sequential-counter encodings use "
    "auxiliary variables to reduce clause growth for larger cardinalities. "
    "The exact encoding and implementation details should be checked against "
    "the source code before submission."
)

doc.add_heading("2.3 CP-SAT and MIP", level=2)
doc.add_paragraph(
    "CP-SAT expresses the Sudoku rules as integer constraints and solves "
    "the resulting constraint model using OR-Tools. The MIP formulations "
    "use binary decision variables x[r,c,v] and linear constraints."
)
doc.add_paragraph(
    "For each cell, the sum over all candidate values equals one. "
    "For each row, column, and subgrid, the sum of variables associated "
    "with any fixed value equals one. Clues are imposed as fixed variables."
)
doc.add_paragraph(
    "Although the formulations express equivalent Sudoku rules, their "
    "internal propagation, presolve, branching, and conflict-learning "
    "mechanisms differ. Runtime comparisons therefore measure complete "
    "solver systems, not just mathematical notation."
)

# 3
doc.add_heading("3. Experimental design", level=1)
doc.add_heading("3.1 Datasets", level=2)
add_table(doc, ["Dataset", "Instances", "Purpose"], [
    ["9x9", "30: 10 easy, 10 medium, 10 hard",
     "Difficulty-based comparison"],
    ["16x16", "10", "Scaled Sudoku evaluation"],
    ["25x25", "10", "Larger-instance evaluation"],
])
doc.add_paragraph(
    "The 9x9 instances are stored in data/sudoku_instances.json. "
    "The scaled datasets are stored under data/scaled-sudoku-sample/. "
    "The scaled-instance format uses '.' for empty cells and symbols "
    "to represent values above 9. Dataset provenance and any transformations "
    "should be documented with the final submission."
)

doc.add_heading("3.2 Solvers", level=2)
add_table(doc, ["Solver", "Model family", "Role"], [
    ["SAT", "Boolean satisfiability", "Boolean encoding of Sudoku rules"],
    ["CP-SAT", "Constraint programming / SAT hybrid",
     "Integer constraint model"],
    ["CPLEX MIP", "Mixed-integer programming",
     "Binary-variable linear model"],
    ["Gurobi", "Mixed-integer programming",
     "Binary-variable linear model"],
    ["CPLEX CP", "Constraint programming",
     "Additional 9x9 comparison where recorded"],
])

doc.add_heading("3.3 Timing and validation", level=2)
doc.add_paragraph(
    "The 9x9 benchmark CSV records five measured runs per instance and "
    "solver, with median_time_seconds as the per-instance statistic. "
    "The scaled benchmark CSV records size, puzzle, solver, status, "
    "validity, and elapsed time. Scaled solver comparisons in this report "
    "include only records marked SOLVED and valid=True."
)
doc.add_paragraph(
    "Timing measurements depend on hardware, operating system, library "
    "versions, solver parameters, warm-up policy, and system load. The "
    "hardware model and exact package versions were not supplied to this "
    "report generator and must be recorded separately for reproducibility."
)

# 4: 9x9 results
doc.add_heading("4. Results on 9x9 Sudoku", level=1)
rows9 = read_csv(
    RESULTS / "benchmark_sat_cpsat_cplex_mip_cplex_cp_gurobi.csv"
)

if rows9:
    groups = {}
    for row in rows9:
        try:
            t = float(row["median_time_seconds"])
        except (ValueError, KeyError):
            continue
        key = (row["difficulty"], row["solver"])
        groups.setdefault(key, []).append(t)

    difficulties = ["easy", "medium", "hard"]
    solvers = sorted({solver for _, solver in groups})
    table_rows = []
    for difficulty in difficulties:
        for solver in solvers:
            values = groups.get((difficulty, solver), [])
            if values:
                table_rows.append([
                    difficulty.title(),
                    solver,
                    str(len(values)),
                    fmt_seconds(statistics.mean(values)),
                    fmt_seconds(statistics.median(values)),
                ])

    add_table(
        doc,
        ["Difficulty", "Solver", "Instances", "Mean (s)", "Median (s)"],
        table_rows,
    )
    doc.add_paragraph(
        "Table values are calculated from the per-instance median times "
        "recorded in the CSV. The mean shown here is the arithmetic mean "
        "of those instance-level medians, not the mean of all individual runs."
    )
else:
    doc.add_paragraph("The 9x9 benchmark CSV was not found.")

add_figure(
    doc,
    RESULTS / "sat_vs_cpsat_cplex_gurobi.png",
    "Figure 1. Existing 9x9 comparison chart. Verify its solver set and "
    "data provenance against the benchmark CSV before final submission.",
)

doc.add_heading("4.1 Interpretation", level=2)
doc.add_paragraph(
    "The 9x9 results allow a comparison across difficulty groups. Small "
    "differences should not be treated as universal performance rankings. "
    "The benchmark covers a finite dataset, and the observed ordering may "
    "change with puzzle selection, solver versions, parameters, and hardware."
)

# 5: scaled results
doc.add_heading("5. Results on scaled Sudoku", level=1)
scaled = read_csv(RESULTS / "benchmark_scaled_all_solvers.csv")
valid_rows = []
for row in scaled:
    try:
        valid = row.get("valid", "").strip().lower() == "true"
        solved = row.get("status", "").strip().upper() == "SOLVED"
        seconds = float(row["time_seconds"])
        size = int(row["size"])
    except (ValueError, KeyError):
        continue
    if valid and solved and seconds >= 0:
        valid_rows.append({
            "size": size,
            "solver": row["solver"],
            "seconds": seconds,
            "puzzle": row.get("puzzle", ""),
        })

if valid_rows:
    table_rows = []
    for size in sorted({r["size"] for r in valid_rows}):
        for solver in sorted({
            r["solver"] for r in valid_rows if r["size"] == size
        }):
            values = [
                r["seconds"] for r in valid_rows
                if r["size"] == size and r["solver"] == solver
            ]
            table_rows.append([
                f"{size}x{size}",
                solver,
                str(len(values)),
                fmt_seconds(statistics.mean(values)),
                fmt_seconds(statistics.median(values)),
            ])

    add_table(
        doc,
        ["Grid", "Solver", "Valid solutions", "Mean (s)", "Median (s)"],
        table_rows,
    )
else:
    doc.add_paragraph(
        "No valid solved records were found in the scaled benchmark CSV."
    )

add_figure(
    doc,
    RESULTS / "scaled_sudoku_16x16_comparison.png",
    "Figure 2. Runtime comparison for 16x16 Sudoku instances.",
)
add_figure(
    doc,
    RESULTS / "scaled_sudoku_25x25_comparison.png",
    "Figure 3. Runtime comparison for 25x25 Sudoku instances.",
)

doc.add_heading("5.1 Interpretation", level=2)
doc.add_paragraph(
    "The scaled results should be interpreted separately from the 9x9 "
    "experiment. Larger grids increase the number of decision variables "
    "and constraints, and runtime can grow nonlinearly. Mean and median "
    "are both reported because a small number of difficult instances can "
    "strongly affect the mean."
)
doc.add_paragraph(
    "Only valid solved records are included in the summary table. This "
    "filter makes the timing comparison easier to interpret but does not "
    "by itself describe timeouts or failures. Any omitted or timed-out "
    "runs should be discussed explicitly in the final evaluation."
)

# 6
doc.add_heading("6. Threats to validity and limitations", level=1)
add_bullet(
    doc, "The benchmark uses a limited number of instances at each grid size."
)
add_bullet(
    doc, "Runtime depends on the computer, operating system, solver versions, "
    "parameters, and background load."
)
add_bullet(
    doc, "Different solvers may use different default heuristics and presolve "
    "strategies, so this is a system-level comparison."
)
add_bullet(
    doc, "Only successful, valid scaled solutions appear in the scaled summary; "
    "timeouts and errors require separate reporting."
)
add_bullet(
    doc, "The report generator does not independently verify every solution "
    "grid; it relies on the benchmark's recorded validity field."
)
add_bullet(
    doc, "The exact hardware and software version inventory must be added "
    "before presenting the experiment as fully reproducible."
)

# 7
doc.add_heading("7. Conclusion and future work", level=1)
doc.add_paragraph(
    "This project provides a practical comparison of SAT, CP-SAT, and MIP "
    "approaches to Sudoku, with additional results for larger grids. The "
    "stored CSV files make it possible to regenerate summary tables and "
    "compare valid solving times consistently. No single solver should be "
    "declared universally superior from these experiments alone."
)
doc.add_paragraph(
    "Future work includes increasing the number and diversity of instances, "
    "recording hardware and package versions, standardizing time limits and "
    "warm-up procedures, reporting timeout rates, and evaluating SAT "
    "cardinality encodings using clause and variable counts as well as time."
)

# References
doc.add_heading("References", level=1)
references = [
    "Biere, A., Heule, M., van Maaren, H., and Walsh, T. (eds.). "
    "Handbook of Satisfiability. IOS Press, 2009.",
    "Google OR-Tools. CP-SAT Solver documentation. "
    "https://developers.google.com/optimization/cp/cp_solver",
    "IBM. IBM ILOG CPLEX Optimization Studio documentation. "
    "https://www.ibm.com/docs/en/icos",
    "Gurobi Optimization. Gurobi Optimizer documentation. "
    "https://docs.gurobi.com/",
    "Scaled Sudoku Instances dataset. "
    "https://github.com/zayenz/scaled-sudoku-instances",
    "Project repository: Sudoku SAT Solver. "
    "https://github.com/transonnn/sudoku-sat-solver",
]
for reference in references:
    doc.add_paragraph(reference, style="List Number")

# Appendix
doc.add_heading("Appendix A. Reproducing the experiments", level=1)
doc.add_paragraph(
    "Run commands from the repository root after activating the project "
    "virtual environment. The exact command and required solver installation "
    "may depend on the current scripts and license configuration."
)
for command in [
    "python -m pytest -v",
    "python -m benchmarks.benchmark",
    "python -m benchmarks.benchmark_scaled_solvers",
    "python -m benchmarks.plot_scaled_results",
]:
    p = doc.add_paragraph()
    r = p.add_run(command)
    r.font.name = "Consolas"
    r.font.size = Pt(9)

doc.add_paragraph(
    "Before submission, confirm that the command names match the current "
    "repository, update the author and institution fields, record the "
    "hardware and software versions, and verify that every figure caption "
    "matches the data used to generate the figure."
)

# Footer with page number
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Sudoku Solver Comparative Study | Page ")

    add_word_field(footer, "PAGE")

doc.save(OUTPUT)
print(f"Report created: {OUTPUT}")
print(f"9x9 benchmark records: {len(rows9)}")
print(f"Scaled benchmark records: {len(scaled)}")
print(f"Valid solved scaled records included: {len(valid_rows)}")
